
from io import BytesIO

from docx import Document
from fastapi import APIRouter, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse

router = APIRouter(prefix="/api/word", tags=["Word Converter"])

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB


async def read_upload(file: UploadFile, expected_extension: str) -> bytes:
    filename = (file.filename or "").lower()

    if not filename.endswith(expected_extension):
        raise HTTPException(
            status_code=400,
            detail=f"Please upload a {expected_extension} file",
        )

    data = await file.read(MAX_FILE_SIZE + 1)
    await file.close()

    if not data:
        raise HTTPException(status_code=400, detail="The file is empty")

    if len(data) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="Maximum file size is 5 MB")

    return data


@router.post("/docx-to-txt")
async def docx_to_txt(file: UploadFile = File(...)):
    data = await read_upload(file, ".docx")

    try:
        document = Document(BytesIO(data))

        lines = [paragraph.text for paragraph in document.paragraphs]

        # Include table content, row by row.
        for table in document.tables:
            for row in table.rows:
                lines.append("\t".join(cell.text for cell in row.cells))

        text_content = "\n".join(lines)
        output = BytesIO(text_content.encode("utf-8"))
        output.seek(0)

    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Unable to read this Word document",
        )

    return StreamingResponse(
        output,
        media_type="text/plain; charset=utf-8",
        headers={
            "Content-Disposition": 'attachment; filename="filezen-document.txt"',
            "Cache-Control": "no-store",
        },
    )


@router.post("/txt-to-docx")
async def txt_to_docx(file: UploadFile = File(...)):
    data = await read_upload(file, ".txt")

    try:
        text_content = data.decode("utf-8-sig")
        document = Document()

        for line in text_content.splitlines():
            document.add_paragraph(line)

        output = BytesIO()
        document.save(output)
        output.seek(0)

    except UnicodeDecodeError:
        raise HTTPException(
            status_code=400,
            detail="The text file must use UTF-8 encoding",
        )
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Unable to convert this text file",
        )

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": 'attachment; filename="filezen-document.docx"',
            "Cache-Control": "no-store",
        },
    )
