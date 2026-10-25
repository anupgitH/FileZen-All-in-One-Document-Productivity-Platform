import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Literal

import fitz  # PyMuPDF
import pdfplumber
from docx import Document
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from starlette.background import BackgroundTask

router = APIRouter(
    prefix="/api/documents",
    tags=["Document conversions"],
)

MAX_UPLOAD_BYTES = 20 * 1024 * 1024

ALLOWED_EXTENSIONS = {
    "pdf-to-word": {".pdf"},
    "pdf-to-powerpoint": {".pdf"},
    "pdf-to-excel": {".pdf"},
    "word-to-pdf": {".docx"},
    "powerpoint-to-pdf": {".pptx"},
    "excel-to-pdf": {".xlsx"},
    "edit-pdf": {".pdf"},
}


def _check_office():
    """Check whether LibreOffice is installed."""
    executable = shutil.which("libreoffice") or shutil.which("soffice")

    if not executable:
        raise HTTPException(
            status_code=503,
            detail=(
                "LibreOffice is required for this conversion. "
                "Install LibreOffice on the server."
            ),
        )

    return executable


async def _save_upload(
    upload: UploadFile,
    folder: Path,
    allowed: set[str],
) -> Path:
    """Validate and save an uploaded file with a size limit."""

    name = Path(upload.filename or "upload").name
    extension = Path(name).suffix.lower()

    if extension not in allowed:
        raise HTTPException(
            status_code=415,
            detail=(
                "Unsupported file extension. Allowed: "
                + ", ".join(sorted(allowed))
            ),
        )

    destination = folder / f"input{extension}"
    size = 0

    try:
        with destination.open("wb") as output:
            while True:
                chunk = await upload.read(1024 * 1024)

                if not chunk:
                    break

                size += len(chunk)

                if size > MAX_UPLOAD_BYTES:
                    raise HTTPException(
                        status_code=413,
                        detail="File exceeds the 20 MB upload limit.",
                    )

                output.write(chunk)

    finally:
        await upload.close()

    if size == 0:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty.",
        )

    return destination


def _persistent_response_copy(
    source: Path,
    filename: str,
    media_type: str,
):
    """
    Copy the result outside the temporary working directory.
    Delete the copy after FastAPI finishes sending the response.
    """

    fd, path = tempfile.mkstemp(
        prefix="filezen-result-",
        suffix=source.suffix,
    )

    os.close(fd)

    try:
        shutil.copyfile(source, path)
    except Exception:
        if os.path.exists(path):
            os.remove(path)
        raise

    return FileResponse(
        path=path,
        filename=filename,
        media_type=media_type,
        background=BackgroundTask(
            lambda: os.path.exists(path) and os.remove(path)
        ),
    )


def _run_libreoffice(
    executable: str,
    source: Path,
    output_dir: Path,
    output_format: str,
):
    """Convert a document using LibreOffice in headless mode."""

    profile = output_dir / "lo-profile"

    command = [
        executable,
        f"-env:UserInstallation={profile.as_uri()}",
        "--headless",
        "--convert-to",
        output_format,
        "--outdir",
        str(output_dir),
        str(source),
    ]

    try:
        result = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=90,
            check=False,
        )

    except subprocess.TimeoutExpired as exc:
        raise HTTPException(
            status_code=504,
            detail="Conversion timed out.",
        ) from exc

    output = output_dir / f"{source.stem}.{output_format}"

    if result.returncode != 0 or not output.exists():
        detail = (
            result.stderr
            or result.stdout
            or "LibreOffice could not convert this file."
        )[-1000:]

        raise HTTPException(
            status_code=422,
            detail=detail,
        )

    return output


# --------------------------------------------------
# 1. PDF TO WORD
# POST /api/documents/pdf-to-word
# --------------------------------------------------

@router.post("/pdf-to-word")
async def pdf_to_word(file: UploadFile = File(...)):
    with tempfile.TemporaryDirectory(
        prefix="filezen-pdf-word-"
    ) as tmp:

        folder = Path(tmp)

        source = await _save_upload(
            file,
            folder,
            ALLOWED_EXTENSIONS["pdf-to-word"],
        )

        target = folder / "output.docx"

        try:
            from pdf2docx import Converter

            converter = Converter(str(source))

            try:
                converter.convert(
                    str(target),
                    start=0,
                    end=None,
                )
            finally:
                converter.close()

        except HTTPException:
            raise

        except Exception as exc:
            raise HTTPException(
                status_code=422,
                detail=f"PDF-to-Word conversion failed: {exc}",
            ) from exc

        if not target.exists():
            raise HTTPException(
                status_code=422,
                detail="Could not create the Word document.",
            )

        return _persistent_response_copy(
            target,
            "filezen-converted.docx",
            (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )


# --------------------------------------------------
# 2. PDF TO POWERPOINT
# POST /api/documents/pdf-to-powerpoint
# --------------------------------------------------

@router.post("/pdf-to-powerpoint")
async def pdf_to_powerpoint(file: UploadFile = File(...)):
    with tempfile.TemporaryDirectory(
        prefix="filezen-pdf-pptx-"
    ) as tmp:

        folder = Path(tmp)

        source = await _save_upload(
            file,
            folder,
            ALLOWED_EXTENSIONS["pdf-to-powerpoint"],
        )

        target = folder / "output.pptx"

        try:
            pdf = fitz.open(source)

            if len(pdf) > 100:
                pdf.close()

                raise HTTPException(
                    status_code=413,
                    detail="PDF has too many pages (maximum 100).",
                )

            presentation = Presentation()

            presentation.slide_width = Inches(10)
            presentation.slide_height = Inches(7.5)

            blank_layout = presentation.slide_layouts[6]

            try:
                for page in pdf:
                    pixmap = page.get_pixmap(
                        matrix=fitz.Matrix(1.5, 1.5),
                        alpha=False,
                    )

                    image_path = (
                        folder / f"page-{page.number + 1}.png"
                    )

                    pixmap.save(str(image_path))

                    slide = presentation.slides.add_slide(
                        blank_layout
                    )

                    slide.shapes.add_picture(
                        str(image_path),
                        0,
                        0,
                        width=presentation.slide_width,
                        height=presentation.slide_height,
                    )

            finally:
                pdf.close()

            presentation.save(str(target))

        except HTTPException:
            raise

        except Exception as exc:
            raise HTTPException(
                status_code=422,
                detail=(
                    "PDF-to-PowerPoint conversion failed: "
                    f"{exc}"
                ),
            ) from exc

        return _persistent_response_copy(
            target,
            "filezen-converted.pptx",
            (
                "application/vnd.openxmlformats-officedocument."
                "presentationml.presentation"
            ),
        )


# --------------------------------------------------
# 3. PDF TO EXCEL
# POST /api/documents/pdf-to-excel
# --------------------------------------------------

@router.post("/pdf-to-excel")
async def pdf_to_excel(file: UploadFile = File(...)):
    with tempfile.TemporaryDirectory(
        prefix="filezen-pdf-xlsx-"
    ) as tmp:

        folder = Path(tmp)

        source = await _save_upload(
            file,
            folder,
            ALLOWED_EXTENSIONS["pdf-to-excel"],
        )

        target = folder / "output.xlsx"

        try:
            workbook = Workbook()

            # Remove the default worksheet.
            default_sheet = workbook.active
            workbook.remove(default_sheet)

            with pdfplumber.open(source) as pdf:

                if len(pdf.pages) > 200:
                    raise HTTPException(
                        status_code=413,
                        detail=(
                            "PDF has too many pages "
                            "(maximum 200)."
                        ),
                    )

                for page_index, page in enumerate(
                    pdf.pages,
                    start=1,
                ):

                    tables = page.extract_tables() or []

                    for table_index, table in enumerate(
                        tables,
                        start=1,
                    ):

                        sheet_name = (
                            f"P{page_index}T{table_index}"
                        )[:31]

                        sheet = workbook.create_sheet(
                            sheet_name
                        )

                        for row in table:
                            sheet.append([
                                cell if cell is not None else ""
                                for cell in row
                            ])

            # If no tables were detected, export extracted text.
            if not workbook.sheetnames:
                sheet = workbook.create_sheet(
                    "Extracted Text"
                )

                with pdfplumber.open(source) as pdf:
                    for page_index, page in enumerate(
                        pdf.pages,
                        start=1,
                    ):

                        sheet.append([
                            f"Page {page_index}"
                        ])

                        extracted_text = (
                            page.extract_text() or ""
                        )

                        for line in extracted_text.splitlines():
                            sheet.append([line])

            workbook.save(str(target))

        except HTTPException:
            raise

        except Exception as exc:
            raise HTTPException(
                status_code=422,
                detail=(
                    "PDF-to-Excel extraction failed: "
                    f"{exc}"
                ),
            ) from exc

        return _persistent_response_copy(
            target,
            "filezen-extracted.xlsx",
            (
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )


# --------------------------------------------------
# 4. WORD TO PDF
# POST /api/documents/word-to-pdf
# Accepts DOCX files
# --------------------------------------------------

@router.post("/word-to-pdf")
async def word_to_pdf(file: UploadFile = File(...)):
    executable = _check_office()

    with tempfile.TemporaryDirectory(
        prefix="filezen-docx-pdf-"
    ) as tmp:

        folder = Path(tmp)

        source = await _save_upload(
            file,
            folder,
            ALLOWED_EXTENSIONS["word-to-pdf"],
        )

        target = _run_libreoffice(
            executable,
            source,
            folder,
            "pdf",
        )

        return _persistent_response_copy(
            target,
            "filezen-converted.pdf",
            "application/pdf",
        )


# --------------------------------------------------
# 5. POWERPOINT TO PDF
# POST /api/documents/powerpoint-to-pdf
# Accepts PPTX files
# --------------------------------------------------

@router.post("/powerpoint-to-pdf")
async def powerpoint_to_pdf(file: UploadFile = File(...)):
    executable = _check_office()

    with tempfile.TemporaryDirectory(
        prefix="filezen-pptx-pdf-"
    ) as tmp:

        folder = Path(tmp)

        source = await _save_upload(
            file,
            folder,
            ALLOWED_EXTENSIONS["powerpoint-to-pdf"],
        )

        target = _run_libreoffice(
            executable,
            source,
            folder,
            "pdf",
        )

        return _persistent_response_copy(
            target,
            "filezen-converted.pdf",
            "application/pdf",
        )


# --------------------------------------------------
# 6. EXCEL TO PDF
# POST /api/documents/excel-to-pdf
# Accepts XLSX files
# --------------------------------------------------

@router.post("/excel-to-pdf")
async def excel_to_pdf(file: UploadFile = File(...)):
    executable = _check_office()

    with tempfile.TemporaryDirectory(
        prefix="filezen-xlsx-pdf-"
    ) as tmp:

        folder = Path(tmp)

        source = await _save_upload(
            file,
            folder,
            ALLOWED_EXTENSIONS["excel-to-pdf"],
        )

        target = _run_libreoffice(
            executable,
            source,
            folder,
            "pdf",
        )

        return _persistent_response_copy(
            target,
            "filezen-converted.pdf",
            "application/pdf",
        )


# --------------------------------------------------
# 7. BASIC PDF EDITOR
# POST /api/documents/edit-pdf
#
# Actions:
#   add-text
#   highlight
#
# Coordinates are in PDF points, from the top-left.
# --------------------------------------------------

@router.post("/edit-pdf")
async def edit_pdf(
    file: UploadFile = File(...),
    action: Literal["add-text", "highlight"] = Form(...),
    page_number: int = Form(1),
    text: str = Form(""),
    x: float = Form(50),
    y: float = Form(50),
    width: float = Form(200),
    height: float = Form(30),
    font_size: float = Form(12),
):
    with tempfile.TemporaryDirectory(
        prefix="filezen-edit-pdf-"
    ) as tmp:

        folder = Path(tmp)

        source = await _save_upload(
            file,
            folder,
            ALLOWED_EXTENSIONS["edit-pdf"],
        )

        if page_number < 1:
            raise HTTPException(
                status_code=400,
                detail="Page number starts at 1.",
            )

        if action == "add-text" and not text.strip():
            raise HTTPException(
                status_code=400,
                detail="Enter text to add.",
            )

        if width <= 0 or height <= 0 or font_size <= 0:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Width, height, and font size "
                    "must be positive."
                ),
            )

        target = folder / "edited.pdf"

        try:
            pdf = fitz.open(source)

            try:
                if page_number > len(pdf):
                    raise HTTPException(
                        status_code=400,
                        detail=(
                            f"PDF has only {len(pdf)} pages."
                        ),
                    )

                page = pdf[page_number - 1]

                rectangle = fitz.Rect(
                    x,
                    y,
                    x + width,
                    y + height,
                )

                # Keep the editing area within page bounds.
                rectangle = rectangle & page.rect

                if rectangle.is_empty:
                    raise HTTPException(
                        status_code=400,
                        detail=(
                            "The selected edit area is "
                            "outside the page."
                        ),
                    )

                if action == "add-text":
                    page.insert_text(
                        fitz.Point(
                            rectangle.x0,
                            rectangle.y0
                            + min(font_size, rectangle.height),
                        ),
                        text[:2000],
                        fontsize=min(font_size, 72),
                        color=(0.1, 0.1, 0.1),
                    )

                elif action == "highlight":
                    annotation = page.add_highlight_annot(
                        rectangle
                    )
                    annotation.update()

                pdf.save(
                    str(target),
                    garbage=4,
                    deflate=True,
                )

            finally:
                pdf.close()

        except HTTPException:
            raise

        except Exception as exc:
            raise HTTPException(
                status_code=422,
                detail=f"PDF edit failed: {exc}",
            ) from exc

        return _persistent_response_copy(
            target,
            "filezen-edited.pdf",
            "application/pdf",
        )